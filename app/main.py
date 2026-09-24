import os
import time
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document
import json

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("GROQ_API_KEY environment variable is not set.")

client=Groq(api_key=my_api_key)

model = "openai/gpt-oss-120b"
# role="user"

job_description=""" 
    Do you want to solve real customer problems through innovative technology? Do you enjoy working on scalable services in a collaborative team environment? Do you want to see your code directly impact millions of customers worldwide?
    
    Key job responsibilities
• Collaborate and communicate effectively with experienced cross-disciplinary Amazonians to design, build, and operate innovative products and services that delight our customers, while participating in technical discussions to drive solutions forward.
• Design and develop scalable solutions using cloud-native architectures and microservices in a large distributed computing environment.
• Participate in code reviews and contribute to technical documentation.
• Build and maintain resilient distributed systems that are scalable, fault-tolerant, and cost-effective.
• Leverage and contribute to the development of GenAI and AI-powered tools to enhance development productivity while staying current with emerging technologies.
• Write clean, maintainable code following best practices and design patterns.
• Work in an agile environment practicing CI/CD principles while participating in operational responsibilities including on-call duties.
• Demonstrate operational excellence through monitoring, troubleshooting, and resolving production issues.

Basic Qualifications
- Experience with at least one general-purpose programming language such as Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structure implementation, basic algorithm development, and/or object-oriented design principles
- Currently has, or is in the process of obtaining a bachelor’s degree in Computer Science, Computer Engineering, Data Science, Information Systems, or related STEM fields
- Must be 18 years of age of older
Preferred Qualifications
- Experience from previous technical internship(s) or demonstrated project experience
- Experience with one or more of the following: AI tools for development productivity, Cloud platforms (preferably AWS), Database systems (SQL and NoSQL), Contributing to open-source projects, Version control systems, Debugging and troubleshooting complex systems
- Demonstrated ability to learn and adapt to new technologies quickly
- Basic understanding of software development lifecycle (SDLC)
- Strong problem-solving and analytical skills
- Excellent written and verbal communication skills
"""

class jobD(BaseModel):
    role: str | None = None
    required_skills: list[str] = []
    preferred_skills: list[str] = []
    minimum_experience: float | None = None
    education_requirements: list[str] = []
    responsibilities: list[str] = []
    
jobd_schema = jobD.model_json_schema()

system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze job descriptions and extract structured information from them. Return only valid JSON matching this schema: {jobd_schema}.

IMPORTANT:
- If the role title is not explicitly stated in the text, infer the most suitable title (e.g., "Software Development Engineer") from the responsibilities.
- Fill the schema with actual information extracted from the job description.
- If minimum experience is not explicitly mentioned, return null for that field.
- If information for a list is missing, return an empty list.
- Do not invent skills or qualifications not mentioned or implied by the job description.
"""

user_prompt = f"""
Analyze the following job description {job_description}
"""
message_system={
    "role": "system",
    "content": system_prompt
}

message_user={
    "role": "user",
    "content": user_prompt
}

response_format={
    "type": "json_object",
}

messages = [message_system, message_user]

response = client.chat.completions.create(
    model=model,
    messages=messages,
    response_format=response_format
)

answer = response.choices[0].message.content
job_data = json.loads(answer)
job = jobD(**job_data)

print(f"Role: {job.role}")
print(f"Minimum Experience: {job.minimum_experience}")
print(f"Education Requirements: {job.education_requirements}")


# Parse real

class MatchResult(BaseModel):
    score: float
    details: dict
    
class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []
    
class Resume(BaseModel):
    name:str | None = None
    email:str | None = None
    phone:str | None = None
    
    total_experience: float | None = None
    
    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []
    
resume_schema = Resume.model_json_schema()
def final_score(job, resume):
    match_schema = MatchResult.model_json_schema()
    
    prompt = f"""
    You are an expert HR recruiter.
    Compare the candidate's resume with the job description.
    
    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}
    
    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    
    Return JSON strictly matching this schema: {match_schema}
    
    Schema requirements:
    - "score": numeric match percentage from 0 to 100 representing overall candidate fit.
    - "details": a JSON object with:
        - "candidate_name": Candidate name
        - "matching_skills": List of matching skills
        - "missing_skills": List of missing important skills
        - "experience_requirement_met": Whether candidate meets experience criteria
        - "overall_match_percentage": Score from 0 to 100
        - "final_verdict": Short concise summary verdict
    
    Keep the response concise and strictly adhere to the schema.
    """
    
    message = {
        "role": "user",
        "content": prompt
    }
    
    messages = [message]
    
    response_format = {
        "type": "json_object",
    }
    
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)

def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.
    
    Extract information from the resume based on its meaning, not only based on exact section headings.
    Different resumes may use different headings
    
    for example:
    - Experience
    - Professional Experience
    - work history
    - Employment
    - Internships
    
    these may all contains relevent experience.
    Skills may also appear in the skills section, work experience , internships or projects.
    
    Return ONLY valid JSON matching this schema: {resume_schema}
    
    Important rules:
    
    1. Do Not invent information.
    2. If a value is not available, returm null.
    3. If a list have no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills from all sections of the resume, not just the skills section.
    
    """

    user_prompt = f"""
        parse the following resume {resume_text}
    """

    message_system = {
        "role": "system",
        "content": system_prompt
    }
    
    message_user = {
        "role": "user", 
        "content": user_prompt
    }
    
    messages = [message_system, message_user]
    
    response_format = {
        "type": "json_object"
    }
    
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    path = Path(file_path)
    if path.suffix.lower() == ".pdf":
        return read_pdf(path)
    elif path.suffix.lower() == ".docx":
        return read_docx(path)
    else:
        raise ValueError(f"Unsupported file format '{path.suffix}'. Please provide a PDF or DOCX file.")


if __name__ == "__main__":
    resume_folder = Path(__file__).resolve().parent / "resumes"
    if not resume_folder.exists():
        resume_folder = Path("resumes")

    all_results = []
    resume_files = [f for f in resume_folder.iterdir() if f.suffix.lower() in [".pdf", ".docx"]]
    print(f"\nFound {len(resume_files)} resume(s) in {resume_folder}")

    for file_path in resume_files:
        print(f"\nProcessing file: {file_path.name}")
        resume_text = read_resume(file_path)
        print("  Parsing resume with LLM...")
        parsed_resume = parse_resume(resume_text)
        print(f"  Candidate Name: {parsed_resume.name}")

        time.sleep(2)
        print("  Evaluating match against Job Description...")
        result = final_score(job, parsed_resume)
        time.sleep(2)
        print(f"  Match Score: {result.score}/100")

        all_results.append({
            "name": parsed_resume.name or file_path.name,
            "file": file_path.name,
            "score": result.score,
            "details": result.details
        })

    all_results.sort(key=lambda candidate: candidate["score"], reverse=True)

    print("\n" + "=" * 60)
    print("ALL CANDIDATE RESULTS (RANKED)")
    print("=" * 60)
    for rank, candidate in enumerate(all_results, 1):
        print(f"\nRank #{rank}: {candidate['name']} ({candidate['file']}) - Score: {candidate['score']}/100")
        print("Details:")
        print(json.dumps(candidate["details"], indent=2))

    if len(all_results) > 2:
        top_2 = all_results[:2]
        worst_2 = all_results[-2:]
        print("\n" + "=" * 60)
        print("TOP 2 CANDIDATES")
        print("=" * 60)
        for candidate in top_2:
            print(f"- {candidate['name']} ({candidate['file']}): {candidate['score']}/100")

        print("\n" + "=" * 60)
        print("LOWEST 2 CANDIDATES")
        print("=" * 60)
        for candidate in worst_2:
            print(f"- {candidate['name']} ({candidate['file']}): {candidate['score']}/100")